"""
End-to-end pipeline. Rerun on a fresh monthly export with no manual edits:

    python run_pipeline.py --inputs <dir> --outputs <dir>

Produces the tidy spine (CSV + Parquet), the exceptions file, the data quality
report, the reliability workbook and the corrected timetable workbooks.
"""

from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import analysis as an
import transport_spine as ts

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name=FONT, size=10)
CLASS_FILL = {
    "SYSTEMATIC": PatternFill("solid", fgColor="FCE4D6"),
    "VOLATILE": PatternFill("solid", fgColor="FFF2CC"),
    "OK": PatternFill("solid", fgColor="E2EFDA"),
    "REVIEW": PatternFill("solid", fgColor="F2F2F2"),
    "LOW_DATA": PatternFill("solid", fgColor="D9D9D9"),
}

FILE_PATTERNS = {
    "routes": "Route_List_Area_Wise*.xlsx",
    "roster": "Route_Wise_Student_list*.xlsx",
    "logs": "Transport_RFID_Attendance_Logs*.xlsx",
}


def resolve_inputs(inputs: Path) -> dict[str, Path]:
    found = {}
    for key, pattern in FILE_PATTERNS.items():
        matches = sorted(inputs.glob(pattern))
        if not matches:
            raise FileNotFoundError(f"No file matching {pattern!r} in {inputs}")
        found[key] = matches[0]
    return found


def style_sheet(ws, freeze="A2", widths=None):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
    ws.freeze_panes = freeze
    for i, col in enumerate(ws.iter_cols(), start=1):
        letter = get_column_letter(i)
        if widths and letter in widths:
            ws.column_dimensions[letter].width = widths[letter]
        else:
            longest = max(
                (len(str(c.value)) for c in col if c.value is not None), default=8
            )
            ws.column_dimensions[letter].width = min(max(longest + 2, 10), 52)
    ws.row_dimensions[1].height = 30


def shade_classification(ws, header_name="Classification"):
    headers = [c.value for c in ws[1]]
    if header_name not in headers:
        return
    idx = headers.index(header_name) + 1
    for row in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
        for cell in row:
            fill = CLASS_FILL.get(str(cell.value))
            if fill:
                cell.fill = fill


def write_sheet(writer, df, name, round_cols=2):
    out = df.copy()
    for c in out.columns:
        if pd.api.types.is_float_dtype(out[c]):
            out[c] = out[c].round(round_cols)
    out.to_excel(writer, sheet_name=name[:31], index=False)
    return out


def main(inputs: Path, outputs: Path) -> dict:
    outputs.mkdir(parents=True, exist_ok=True)
    paths = resolve_inputs(inputs)

    res = ts.build_all(paths["routes"], paths["roster"], paths["logs"])
    spine, students, exceptions, audit = (
        res["spine"], res["students"], res["exceptions"], res["audit"]
    )

    # ---------------- Phase 0 outputs ----------------
    tidy = spine[ts.SPINE_COLUMNS + ["in_regime", "dominant_route_key", "passes_sanity"]]
    tidy.to_csv(outputs / "transport_spine.csv", index=False)
    try:
        tidy.to_parquet(outputs / "transport_spine.parquet", index=False)
    except Exception as exc:                                   # pragma: no cover
        audit.record("output", "parquet write failed", 0, str(exc))
    exceptions.to_csv(outputs / "exceptions_quarantined.csv", index=False)

    # ---------------- Analyses ----------------
    trace = an.bus_arrival_trace(spine)
    stops = an.stop_deviation(trace)
    stops_legacy = an.stop_deviation_per_student(spine)
    scorecard = an.route_scorecard(spine)
    arrival = an.school_arrival(spine)
    pm_stops = an.pm_deviation(spine)
    b2r = an.bell_to_rolling(spine)
    rides = an.ride_durations(spine)
    stop_load, route_load = an.load_profile(students, spine)
    dow, wk = an.temporal(spine)
    dwell = an.dwell_profile(trace)
    comp_leg, comp_route, comp_date = an.scan_compliance(spine)

    students = students.copy()
    students["scheduled_ride_min"] = (
        students.reporting_time_min - students.promised_pickup_min
    )

    timetable = an.build_timetable(stops, res["routes"], students, trace, spine)
    feasibility = an.feasibility_check(timetable, trace, spine)
    pm_timetable = an.build_pm_timetable(pm_stops, spine)

    # ---------------- Reliability workbook ----------------
    rel = outputs / "reliability_workbook.xlsx"
    with pd.ExcelWriter(rel, engine="openpyxl") as w:
        write_sheet(w, audit.to_frame(), "Audit Log")
        write_sheet(w, comp_leg, "Scan Compliance by Leg")
        write_sheet(w, comp_route, "Scan Compliance by Route")
        write_sheet(w, comp_date, "Scan Compliance by Date")
        write_sheet(w, stops, "AM Stop Deviation")
        write_sheet(w, stops_legacy, "AM Stop (per-student)")
        write_sheet(w, scorecard, "Route Scorecard")
        write_sheet(w, arrival, "School Arrival")
        write_sheet(w, pm_stops, "PM Stop Deviation")
        write_sheet(w, b2r, "PM Bell to Rolling")
        write_sheet(w, rides, "Ride Durations")
        write_sheet(w, stop_load, "Stop Load")
        write_sheet(w, route_load, "Route Load")
        write_sheet(w, dwell, "Stop Dwell Time")
        write_sheet(w, dow, "Day of Week")
        write_sheet(w, wk, "Weekly Trend")
        write_sheet(w, feasibility, "Timetable Feasibility")

    from openpyxl import load_workbook

    wb = load_workbook(rel)
    for ws in wb.worksheets:
        style_sheet(ws)
        shade_classification(ws)
    wb.save(rel)

    # ---------------- Timetable workbooks ----------------
    for label, table, published_col in [
        ("AM", timetable, "Current Published Pickup"),
        ("PM", pm_timetable, "Current Published Drop"),
    ]:
        path = outputs / f"corrected_timetable_{label}.xlsx"
        cols = [
            "Stop", published_col,
            "Observed Median Pickup" if label == "AM" else "Observed Median Drop",
            "Recommended Window", "Classification", "n", "Confidence",
        ]
        with pd.ExcelWriter(path, engine="openpyxl") as w:
            if label == "AM":
                write_sheet(w, feasibility, "Feasibility Summary")
            for rk, g in table.groupby("route_key"):
                if pd.isna(rk):
                    continue
                write_sheet(w, g[cols], f"Route {int(rk[1:])}")
        wb = load_workbook(path)
        for ws in wb.worksheets:
            style_sheet(ws, widths={"A": 46, "G": 34})
            shade_classification(ws)
        wb.save(path)

    return {
        **res, "trace": trace, "stops": stops, "stops_legacy": stops_legacy,
        "scorecard": scorecard, "arrival": arrival, "pm_stops": pm_stops,
        "b2r": b2r, "rides": rides, "stop_load": stop_load,
        "route_load": route_load, "dow": dow, "wk": wk, "dwell": dwell,
        "comp_leg": comp_leg, "comp_route": comp_route, "comp_date": comp_date,
        "timetable": timetable, "pm_timetable": pm_timetable,
        "feasibility": feasibility, "students": students, "outputs": outputs,
    }


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--inputs", default="/mnt/user-data/uploads")
    ap.add_argument("--outputs", default="/mnt/user-data/outputs")
    a = ap.parse_args()
    r = main(Path(a.inputs), Path(a.outputs))
    print(r["audit"])
    print(f"\nWritten to {a.outputs}")
